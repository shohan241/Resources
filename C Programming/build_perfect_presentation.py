import os
import sys
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# --- Color Constants (matching refer.pptx & novonex delta) ---
COLOR_BG_CARD_LEFT   = RGBColor(10, 15, 24)    # #0A0F18
COLOR_BG_CARD_RIGHT  = RGBColor(5, 8, 15)      # #05080F
COLOR_BORDER_LEFT    = RGBColor(30, 41, 59)    # #1E293B
COLOR_BORDER_RIGHT   = RGBColor(51, 65, 85)    # #334155
COLOR_BORDER_CYAN    = RGBColor(0, 176, 240)   # #00B0F0
COLOR_BORDER_BLUE    = RGBColor(56, 189, 248)  # #38BDF8
COLOR_BORDER_AMBER   = RGBColor(245, 158, 11)  # #F59E0B
COLOR_BORDER_RED     = RGBColor(239, 68, 68)   # #EF4444

COLOR_CYAN           = RGBColor(0, 176, 240)   # #00B0F0
COLOR_BLUE           = RGBColor(0, 178, 255)   # #00B2FF
COLOR_WHITE          = RGBColor(255, 255, 255) # #FFFFFF
COLOR_LIGHT_SLATE    = RGBColor(226, 232, 240) # #E2E8F0
COLOR_MUTED_SLATE    = RGBColor(148, 163, 184) # #94A3B8
COLOR_YELLOW         = RGBColor(252, 211, 77)  # #FCD34D
COLOR_GOLD_DIVIDER   = RGBColor(255, 192, 0)   # #FFC000
COLOR_RED            = RGBColor(209, 37, 36)   # #D12524

COLOR_BG_BOX         = RGBColor(15, 23, 42)    # #0F172A
COLOR_BG_CALLOUT     = RGBColor(17, 24, 39)    # #111827

FONT_HEADING = 'Amasis MT Pro Light'
FONT_CONTENT = 'Calibri'
FONT_CODE    = 'Consolas'
FONT_BRAND   = 'Amasis MT Pro'

LOGO_PATH = 'novonex_logo.png'
OUTPUT_FILE = 'Programming and Problem Solving.pptx'

def add_top_bar(slide, badge_text):
    # Outer white box
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 411480, 192024, 11365992, 521207)
    rect.fill.background()
    rect.line.color.rgb = COLOR_WHITE
    rect.line.width = Pt(1)

    # Left text: "Programming and Problem Solving"
    tb = slide.shapes.add_textbox(548640, 219456, 7000000, 457200)
    tf = tb.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    r1 = p.add_run()
    r1.text = 'Programming'
    r1.font.name = FONT_HEADING
    r1.font.size = Pt(18)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_CYAN

    r2 = p.add_run()
    r2.text = ' and Problem Solving'
    r2.font.name = FONT_HEADING
    r2.font.size = Pt(18)
    r2.font.bold = True
    r2.font.color.rgb = COLOR_WHITE

    # Right badge
    badge = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 7781544, 256032, 3840480, 384048)
    badge.fill.background()
    badge.line.color.rgb = COLOR_CYAN
    badge.line.width = Pt(1)
    pb = badge.text_frame.paragraphs[0]
    pb.text = badge_text
    pb.font.name = FONT_CONTENT
    pb.font.size = Pt(12)
    pb.font.bold = True
    pb.font.color.rgb = COLOR_WHITE
    pb.alignment = PP_ALIGN.CENTER

def add_bottom_logo(slide):
    if os.path.exists(LOGO_PATH):
        slide.shapes.add_picture(LOGO_PATH, 5724144, 6263640, 749808, 521207)

def create_title_slide(prs):
    slide_layout = prs.slide_layouts[6]
    s = prs.slides.add_slide(slide_layout)

    # Title: Programming and Problem Solving
    tb_title = s.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.33), Inches(0.9))
    tf = tb_title.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER

    r1 = p.add_run()
    r1.text = 'P'
    r1.font.name = FONT_HEADING
    r1.font.size = Pt(40)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_CYAN

    r2 = p.add_run()
    r2.text = 'rogramming and Problem Solving'
    r2.font.name = FONT_HEADING
    r2.font.size = Pt(28)
    r2.font.bold = True
    r2.font.color.rgb = COLOR_WHITE

    # Horizontal divider line
    line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3.8), Inches(3.15), Inches(5.73), Pt(2))
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_WHITE
    line.line.color.rgb = COLOR_WHITE

    # Subtitle: NovoNex Delta
    tb_sub = s.shapes.add_textbox(Inches(3.5), Inches(3.3), Inches(6.33), Inches(0.7))
    tf_sub = tb_sub.text_frame
    p_sub = tf_sub.paragraphs[0]
    p_sub.alignment = PP_ALIGN.CENTER

    runs_brand = [
        ('N', Pt(32), COLOR_BLUE),
        ('ovo', Pt(24), COLOR_BLUE),
        ('Nex', Pt(24), COLOR_RED),
        ('  ', Pt(24), COLOR_WHITE),
        ('D', Pt(32), COLOR_BLUE),
        ('el', Pt(24), COLOR_BLUE),
        ('ta', Pt(24), COLOR_RED)
    ]
    for txt, sz, clr in runs_brand:
        r = p_sub.add_run()
        r.text = txt
        r.font.name = FONT_BRAND
        r.font.size = sz
        r.font.color.rgb = clr

    # Bottom logo
    if os.path.exists(LOGO_PATH):
        s.shapes.add_picture(LOGO_PATH, Inches(5.9), Inches(5.4), Inches(1.5), Inches(0.95))

    return s

def create_agenda_slide(prs, topic_num, topic_title, subtopics):
    slide_layout = prs.slide_layouts[6]
    s = prs.slides.add_slide(slide_layout)

    # Centered Agenda Card
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.7), Inches(1.15), Inches(7.93), Inches(4.9))
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_BG_CARD_LEFT
    card.line.color.rgb = COLOR_BORDER_LEFT
    card.line.width = Pt(1.5)

    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(1.2)
    tf.margin_right = Inches(1.2)
    tf.margin_top = Inches(0.35)

    # Pill tag
    p_tag = tf.paragraphs[0]
    p_tag.text = f'TOPIC {topic_num:02d} • AGENDA'
    p_tag.font.name = FONT_CONTENT
    p_tag.font.size = Pt(11)
    p_tag.font.bold = True
    p_tag.font.color.rgb = COLOR_CYAN
    p_tag.alignment = PP_ALIGN.CENTER
    p_tag.space_after = Pt(4)

    # Heading (strictly font size: 18)
    p_head = tf.add_paragraph()
    p_head.text = topic_title
    p_head.font.name = FONT_HEADING
    p_head.font.size = Pt(18)
    p_head.font.bold = True
    p_head.font.color.rgb = COLOR_WHITE
    p_head.alignment = PP_ALIGN.CENTER
    p_head.space_after = Pt(4)

    # Divider line
    p_div = tf.add_paragraph()
    p_div.text = '-------------------------------------------------------'
    p_div.font.name = FONT_CONTENT
    p_div.font.size = Pt(11)
    p_div.font.color.rgb = COLOR_GOLD_DIVIDER
    p_div.alignment = PP_ALIGN.CENTER
    p_div.space_after = Pt(14)

    # Bullet list container (strictly ONE WORD per bullet point, Content font size: 16)
    for st in subtopics:
        words = st.strip().split()
        if len(words) != 1:
            raise ValueError(f"Subtopic '{st}' must be strictly one word! Found {len(words)} words.")
        p = tf.add_paragraph()
        p.space_after = Pt(9)
        p.alignment = PP_ALIGN.LEFT

        r_bullet = p.add_run()
        r_bullet.text = '•   '
        r_bullet.font.name = FONT_CONTENT
        r_bullet.font.size = Pt(16)
        r_bullet.font.bold = True
        r_bullet.font.color.rgb = COLOR_CYAN

        r_text = p.add_run()
        r_text.text = st.strip()
        r_text.font.name = FONT_CONTENT
        r_text.font.size = Pt(16)
        r_text.font.bold = True
        r_text.font.color.rgb = COLOR_LIGHT_SLATE

    add_bottom_logo(s)
    return s

def create_content_slide(prs, badge_text, heading_text, bullets_data, callout_data, right_board_data, code_snippet=None):
    slide_layout = prs.slide_layouts[6]
    s = prs.slides.add_slide(slide_layout)

    add_top_bar(s, badge_text)
    add_bottom_logo(s)

    # Left Panel Card
    card_l = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 411480, 868680, 5577840, 5212080)
    card_l.fill.solid()
    card_l.fill.fore_color.rgb = COLOR_BG_CARD_LEFT
    card_l.line.color.rgb = COLOR_BORDER_LEFT
    card_l.line.width = Pt(1)

    # Left Heading (strictly font size: 18)
    tb_title = s.shapes.add_textbox(548640, 960120, 5303520, 420000)
    tf_t = tb_title.text_frame
    tf_t.word_wrap = True
    pt = tf_t.paragraphs[0]
    rt = pt.add_run()
    rt.text = heading_text
    rt.font.name = FONT_HEADING
    rt.font.size = Pt(18)
    rt.font.bold = True
    rt.font.color.rgb = COLOR_CYAN

    # Left Bullets (strictly font size: 16)
    if code_snippet:
        num_bullets = len(bullets_data)
        bullets_top = 1440000
        bullets_h = 1650000
        tb_bullets = s.shapes.add_textbox(548640, bullets_top, 5303520, bullets_h)
        code_top = 3300000
        code_h = 1950000
    else:
        tb_bullets = s.shapes.add_textbox(548640, 1440000, 5303520, 3900000)

    tf_b = tb_bullets.text_frame
    tf_b.word_wrap = True

    for idx, (kw, expl) in enumerate(bullets_data):
        p = tf_b.paragraphs[0] if idx == 0 else tf_b.add_paragraph()
        p.space_after = Pt(6 if code_snippet else 10)
        
        r_sym = p.add_run()
        r_sym.text = '•  '
        r_sym.font.name = FONT_CONTENT
        r_sym.font.size = Pt(16)
        r_sym.font.bold = True
        r_sym.font.color.rgb = COLOR_YELLOW
        
        r_kw = p.add_run()
        r_kw.text = kw + ': '
        r_kw.font.name = FONT_CONTENT
        r_kw.font.size = Pt(16)
        r_kw.font.bold = True
        r_kw.font.color.rgb = COLOR_YELLOW
        
        r_tx = p.add_run()
        r_tx.text = expl
        r_tx.font.name = FONT_CONTENT
        r_tx.font.size = Pt(16)
        r_tx.font.color.rgb = COLOR_LIGHT_SLATE

    # Left Code Box
    if code_snippet:
        code_box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 594360, code_top, 5212080, code_h)
        code_box.fill.solid()
        code_box.fill.fore_color.rgb = COLOR_BG_BOX
        code_box.line.color.rgb = COLOR_BORDER_BLUE
        code_box.line.width = Pt(1)
        tf_c = code_box.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = Inches(0.2)
        tf_c.margin_top = Inches(0.1)
        
        lines = code_snippet.strip().split('\n')
        for l_idx, line in enumerate(lines):
            p = tf_c.paragraphs[0] if l_idx == 0 else tf_c.add_paragraph()
            p.space_after = Pt(2)
            r = p.add_run()
            r.text = line
            r.font.name = FONT_CODE
            r.font.size = Pt(10)
            if line.strip().startswith('//'):
                r.font.color.rgb = COLOR_MUTED_SLATE
            elif any(k in line for k in ['int ', 'float ', 'double ', 'char ', 'for ', 'while ', 'if ', 'else', 'switch', 'case', 'return', '#include', 'scanf', 'printf']):
                r.font.color.rgb = COLOR_BORDER_BLUE
            else:
                r.font.color.rgb = COLOR_LIGHT_SLATE

    # Left Callout Box at bottom
    c_icon, c_text, c_color = callout_data
    callout = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 594360, 5440680, 5212080, 457200)
    callout.fill.solid()
    callout.fill.fore_color.rgb = COLOR_BG_CALLOUT
    callout.line.color.rgb = c_color
    callout.line.width = Pt(1)
    pc = callout.text_frame.paragraphs[0]
    rc = pc.add_run()
    rc.text = f'{c_icon} {c_text}'
    rc.font.name = FONT_CONTENT
    rc.font.size = Pt(11)
    rc.font.bold = True
    rc.font.color.rgb = c_color
    pc.alignment = PP_ALIGN.CENTER

    # Right Panel Card
    card_r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 6199632, 868680, 5577840, 5212080)
    card_r.fill.solid()
    card_r.fill.fore_color.rgb = COLOR_BG_CARD_RIGHT
    card_r.line.color.rgb = COLOR_BORDER_RIGHT
    card_r.line.width = Pt(1)

    # Right Title
    tb_rt = s.shapes.add_textbox(6309360, 960120, 5349240, 411480)
    prt = tb_rt.text_frame.paragraphs[0]
    r_rt = prt.add_run()
    r_rt.text = right_board_data['title']
    r_rt.font.name = FONT_CONTENT
    r_rt.font.size = Pt(15)
    r_rt.font.bold = True
    r_rt.font.color.rgb = COLOR_CYAN

    # Right Subtitle
    tb_sub = s.shapes.add_textbox(6309360, 1325880, 5349240, 360000)
    psub = tb_sub.text_frame.paragraphs[0]
    r_sub = psub.add_run()
    r_sub.text = right_board_data['subtitle']
    r_sub.font.name = FONT_CONTENT
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = COLOR_MUTED_SLATE

    # Render right board based on type
    b_type = right_board_data.get('type')

    if b_type == 'diagram_box':
        d_box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, 6355080, 1720000, 5266944, 4150000)
        d_box.fill.solid()
        d_box.fill.fore_color.rgb = COLOR_BG_BOX
        d_box.line.color.rgb = right_board_data.get('border_color', COLOR_BORDER_BLUE)
        d_box.line.width = Pt(1)
        tf_d = d_box.text_frame
        tf_d.word_wrap = True
        tf_d.margin_left = Inches(0.25)
        tf_d.margin_top = Inches(0.2)
        
        head = right_board_data.get('header', '')
        if head:
            p0 = tf_d.paragraphs[0]
            p0.text = head + '\n'
            p0.font.name = FONT_CODE
            p0.font.size = Pt(12)
            p0.font.bold = True
            p0.font.color.rgb = COLOR_CYAN

        content = right_board_data.get('content', '')
        for c_idx, line in enumerate(content.split('\n')):
            p = tf_d.add_paragraph() if (head or c_idx > 0) else tf_d.paragraphs[0]
            p.space_after = Pt(2)
            r = p.add_run()
            r.text = line
            r.font.name = FONT_CODE
            r.font.size = Pt(10.5)
            if any(sym in line for sym in ['───', '===', '-->', '┌', '└', '│', '▲', '▼', '|']):
                r.font.color.rgb = COLOR_CYAN
            elif any(w in line for w in ['•', '1.', '2.', '3.', '4.', '5.']):
                r.font.color.rgb = COLOR_YELLOW
            else:
                r.font.color.rgb = COLOR_LIGHT_SLATE

    elif b_type == 'ladder_boxes':
        items = right_board_data.get('items', [])
        start_top = 1750000
        box_h = 470000
        gap = 45000
        for i_idx, item in enumerate(items):
            box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, 6355080, start_top + i_idx * (box_h + gap), 5266944, box_h)
            box.fill.solid()
            box.fill.fore_color.rgb = COLOR_BG_BOX
            box.line.color.rgb = right_board_data.get('border_color', COLOR_BORDER_BLUE)
            box.line.width = Pt(1)
            tf_i = box.text_frame
            tf_i.margin_left = Inches(0.2)
            tf_i.margin_top = Inches(0.08)
            p = tf_i.paragraphs[0]
            r = p.add_run()
            r.text = item
            r.font.name = FONT_CODE
            r.font.size = Pt(11)
            r.font.bold = True
            r.font.color.rgb = COLOR_LIGHT_SLATE

    elif b_type == 'dual_boxes':
        box1_title = right_board_data.get('box1_title', '')
        box1_content = right_board_data.get('box1_content', '')
        box2_title = right_board_data.get('box2_title', '')
        box2_content = right_board_data.get('box2_content', '')

        # Box 1
        b1 = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, 6355080, 1720000, 5266944, 2000000)
        b1.fill.solid()
        b1.fill.fore_color.rgb = COLOR_BG_BOX
        b1.line.color.rgb = COLOR_BORDER_BLUE
        b1.line.width = Pt(1)
        tf1 = b1.text_frame
        tf1.margin_left = Inches(0.2)
        tf1.margin_top = Inches(0.12)
        p1 = tf1.paragraphs[0]
        p1.text = box1_title + '\n'
        p1.font.name = FONT_CODE
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_CYAN
        for line in box1_content.split('\n'):
            p = tf1.add_paragraph()
            p.text = line
            p.font.name = FONT_CODE
            p.font.size = Pt(10)
            p.font.color.rgb = COLOR_LIGHT_SLATE

        # Box 2
        b2 = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, 6355080, 3850000, 5266944, 2000000)
        b2.fill.solid()
        b2.fill.fore_color.rgb = COLOR_BG_BOX
        b2.line.color.rgb = COLOR_BORDER_AMBER
        b2.line.width = Pt(1)
        tf2 = b2.text_frame
        tf2.margin_left = Inches(0.2)
        tf2.margin_top = Inches(0.12)
        p2 = tf2.paragraphs[0]
        p2.text = box2_title + '\n'
        p2.font.name = FONT_CODE
        p2.font.size = Pt(11)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_YELLOW
        for line in box2_content.split('\n'):
            p = tf2.add_paragraph()
            p.text = line
            p.font.name = FONT_CODE
            p.font.size = Pt(10)
            p.font.color.rgb = COLOR_LIGHT_SLATE

    return s

def create_empty_slide(prs):
    slide_layout = prs.slide_layouts[6]
    s = prs.slides.add_slide(slide_layout)
    return s

def create_thank_you_slide(prs):
    slide_layout = prs.slide_layouts[6]
    s = prs.slides.add_slide(slide_layout)

    tb = s.shapes.add_textbox(Inches(3.5), Inches(2.3), Inches(6.33), Inches(0.9))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = 'Thank You'
    r.font.name = FONT_HEADING
    r.font.size = Pt(44)
    r.font.bold = True
    r.font.color.rgb = COLOR_WHITE

    line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.2), Inches(3.3), Inches(4.93), Pt(2))
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_WHITE
    line.line.color.rgb = COLOR_WHITE

    tb_sub = s.shapes.add_textbox(Inches(2.5), Inches(3.45), Inches(8.33), Inches(0.7))
    tf_sub = tb_sub.text_frame
    p_sub = tf_sub.paragraphs[0]
    p_sub.alignment = PP_ALIGN.CENTER
    r_sub = p_sub.add_run()
    r_sub.text = 'Happy Coding with NovoNex Delta!'
    r_sub.font.name = FONT_CONTENT
    r_sub.font.size = Pt(20)
    r_sub.font.bold = True
    r_sub.font.color.rgb = COLOR_CYAN

    if os.path.exists(LOGO_PATH):
        s.shapes.add_picture(LOGO_PATH, Inches(5.8), Inches(5.2), Inches(1.7), Inches(1.1))

    return s

def create_flowchart_shapes_slide(prs):
    slide_layout = prs.slide_layouts[6]
    s = prs.slides.add_slide(slide_layout)

    add_top_bar(s, "Flowchart Shapes")
    add_bottom_logo(s)

    # Left Panel Card
    card_l = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 411480, 868680, 5577840, 5212080)
    card_l.fill.solid()
    card_l.fill.fore_color.rgb = COLOR_BG_CARD_LEFT
    card_l.line.color.rgb = COLOR_BORDER_LEFT
    card_l.line.width = Pt(1)

    # Left Heading
    tb_title_l = s.shapes.add_textbox(548640, 960120, 5303520, 380000)
    tf_tl = tb_title_l.text_frame
    tf_tl.word_wrap = True
    pt_l = tf_tl.paragraphs[0]
    rt_l = pt_l.add_run()
    rt_l.text = "Core Process & I/O Symbols"
    rt_l.font.name = FONT_HEADING
    rt_l.font.size = Pt(18)
    rt_l.font.bold = True
    rt_l.font.color.rgb = COLOR_CYAN

    # Subtitle Left
    tb_sub_l = s.shapes.add_textbox(548640, 1310000, 5303520, 250000)
    psub_l = tb_sub_l.text_frame.paragraphs[0]
    rsub_l = psub_l.add_run()
    rsub_l.text = "[Visual shapes and standard operations for program flow]"
    rsub_l.font.name = FONT_CONTENT
    rsub_l.font.size = Pt(10.5)
    rsub_l.font.color.rgb = COLOR_MUTED_SLATE

    # Right Panel Card
    card_r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 6199632, 868680, 5577840, 5212080)
    card_r.fill.solid()
    card_r.fill.fore_color.rgb = COLOR_BG_CARD_RIGHT
    card_r.line.color.rgb = COLOR_BORDER_RIGHT
    card_r.line.width = Pt(1)

    # Right Heading
    tb_title_r = s.shapes.add_textbox(6309360, 960120, 5349240, 380000)
    tf_tr = tb_title_r.text_frame
    tf_tr.word_wrap = True
    pt_r = tf_tr.paragraphs[0]
    rt_r = pt_r.add_run()
    rt_r.text = "Decision & Routing Symbols"
    rt_r.font.name = FONT_HEADING
    rt_r.font.size = Pt(18)
    rt_r.font.bold = True
    rt_r.font.color.rgb = COLOR_CYAN

    # Subtitle Right
    tb_sub_r = s.shapes.add_textbox(6309360, 1310000, 5349240, 250000)
    psub_r = tb_sub_r.text_frame.paragraphs[0]
    rsub_r = psub_r.add_run()
    rsub_r.text = "[Branching conditions, connector junctions & directional flow]"
    rsub_r.font.name = FONT_CONTENT
    rsub_r.font.size = Pt(10.5)
    rsub_r.font.color.rgb = COLOR_MUTED_SLATE

    # Left Shapes Data
    left_shapes_data = [
        {
            "shape_type": MSO_SHAPE.ROUNDED_RECTANGLE,
            "shape_label": "START / END",
            "font_size": 9.5,
            "is_circle": False,
            "title": "Terminator (Oval / Rounded Pill)",
            "use": "Defines the mandatory entry (Start) and exit (End / Stop) boundary terminals of program flow.",
            "c_code": "C: Entry at main() and exit at return 0;"
        },
        {
            "shape_type": MSO_SHAPE.RECTANGLE,
            "shape_label": "Process",
            "font_size": 10.5,
            "is_circle": False,
            "title": "Process Box (Rectangle)",
            "use": "Represents computational tasks, arithmetic calculations, data transformations, and assignments.",
            "c_code": "C: sum = a + b;  |  count = 0;  |  x++;"
        },
        {
            "shape_type": MSO_SHAPE.PARALLELOGRAM,
            "shape_label": "Input / Output",
            "font_size": 9.5,
            "is_circle": False,
            "title": "Input / Output (Parallelogram)",
            "use": "Represents data receiving operations from user input or printing results to output displays.",
            "c_code": "C: scanf(\"%d\", &a);  |  printf(\"Result: %d\", sum);"
        }
    ]

    # Right Shapes Data
    right_shapes_data = [
        {
            "shape_type": MSO_SHAPE.DIAMOND,
            "shape_label": "Decision?",
            "font_size": 9.5,
            "is_circle": False,
            "title": "Decision Box (Diamond)",
            "use": "Evaluates a boolean conditional question; splits flow into divergent True / False paths.",
            "c_code": "C: if (score >= 50)  |  while (i < n)"
        },
        {
            "shape_type": MSO_SHAPE.OVAL,
            "shape_label": "●",
            "font_size": 14,
            "is_circle": True,
            "title": "Connector (Circle)",
            "use": "Joins multiple intersecting flow lines into a single point to prevent messy crossing lines.",
            "c_code": "C: Merge point of if-else branches, loop returns"
        },
        {
            "shape_type": MSO_SHAPE.RIGHT_ARROW,
            "shape_label": "Flow",
            "font_size": 10.5,
            "is_circle": False,
            "title": "Flow Line / Arrow",
            "use": "Indicates the exact sequence order and directional path of instruction execution.",
            "c_code": "C: Sequential step-by-step instruction execution"
        }
    ]

    start_top = 1600000
    card_h = 1180000
    gap = 85000

    def render_shape_block(parent_slide, x_card, y_card, w_card, h_card, data):
        box = parent_slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_card, y_card, w_card, h_card)
        box.fill.solid()
        box.fill.fore_color.rgb = COLOR_BG_BOX
        box.line.color.rgb = COLOR_BORDER_BLUE
        box.line.width = Pt(1)

        if data.get("is_circle"):
            s_w = Inches(0.85)
            s_h = Inches(0.85)
            s_x = x_card + Inches(0.4)
            s_y = y_card + (h_card - s_h) // 2
        elif data["shape_type"] == MSO_SHAPE.DIAMOND:
            s_w = Inches(1.45)
            s_h = Inches(0.82)
            s_x = x_card + Inches(0.12)
            s_y = y_card + (h_card - s_h) // 2
        elif data["shape_type"] == MSO_SHAPE.RIGHT_ARROW:
            s_w = Inches(1.4)
            s_h = Inches(0.65)
            s_x = x_card + Inches(0.14)
            s_y = y_card + (h_card - s_h) // 2
        else:
            s_w = Inches(1.4)
            s_h = Inches(0.75)
            s_x = x_card + Inches(0.14)
            s_y = y_card + (h_card - s_h) // 2

        visual_shape = parent_slide.shapes.add_shape(data["shape_type"], s_x, s_y, s_w, s_h)
        visual_shape.fill.solid()
        visual_shape.fill.fore_color.rgb = RGBColor(14, 30, 55)
        visual_shape.line.color.rgb = COLOR_CYAN
        visual_shape.line.width = Pt(1.5)

        tf_s = visual_shape.text_frame
        tf_s.word_wrap = False
        tf_s.margin_left = Inches(0.0)
        tf_s.margin_right = Inches(0.0)
        tf_s.margin_top = Inches(0.0)
        tf_s.margin_bottom = Inches(0.0)
        tf_s.vertical_anchor = MSO_ANCHOR.MIDDLE
        ps = tf_s.paragraphs[0]
        ps.alignment = PP_ALIGN.CENTER
        rs = ps.add_run()
        rs.text = data["shape_label"]
        rs.font.name = FONT_CONTENT
        rs.font.size = Pt(data["font_size"])
        rs.font.bold = True
        rs.font.color.rgb = COLOR_WHITE

        desc_x = x_card + Inches(1.68)
        desc_w = (x_card + w_card) - desc_x - Inches(0.12)
        desc_y = y_card + Inches(0.06)
        desc_h = h_card - Inches(0.12)

        tb_desc = parent_slide.shapes.add_textbox(desc_x, desc_y, desc_w, desc_h)
        tf_d = tb_desc.text_frame
        tf_d.word_wrap = True
        tf_d.margin_left = Inches(0.0)
        tf_d.margin_right = Inches(0.0)
        tf_d.margin_top = Inches(0.04)

        p_title = tf_d.paragraphs[0]
        p_title.space_after = Pt(2)
        r_t = p_title.add_run()
        r_t.text = data["title"]
        r_t.font.name = FONT_CONTENT
        r_t.font.size = Pt(12)
        r_t.font.bold = True
        r_t.font.color.rgb = COLOR_YELLOW

        p_use = tf_d.add_paragraph()
        p_use.space_after = Pt(2)
        r_u_lbl = p_use.add_run()
        r_u_lbl.text = "• Purpose: "
        r_u_lbl.font.name = FONT_CONTENT
        r_u_lbl.font.size = Pt(9.5)
        r_u_lbl.font.bold = True
        r_u_lbl.font.color.rgb = COLOR_CYAN

        r_u_txt = p_use.add_run()
        r_u_txt.text = data["use"]
        r_u_txt.font.name = FONT_CONTENT
        r_u_txt.font.size = Pt(9.5)
        r_u_txt.font.color.rgb = COLOR_LIGHT_SLATE

        p_c = tf_d.add_paragraph()
        r_c_lbl = p_c.add_run()
        r_c_lbl.text = "• "
        r_c_lbl.font.name = FONT_CONTENT
        r_c_lbl.font.size = Pt(9)
        r_c_lbl.font.bold = True
        r_c_lbl.font.color.rgb = COLOR_CYAN

        r_c_txt = p_c.add_run()
        r_c_txt.text = data["c_code"]
        r_c_txt.font.name = FONT_CODE
        r_c_txt.font.size = Pt(9)
        r_c_txt.font.bold = True
        r_c_txt.font.color.rgb = COLOR_BORDER_BLUE

    for idx, d in enumerate(left_shapes_data):
        y = start_top + idx * (card_h + gap)
        render_shape_block(s, 548640, y, 5303520, card_h, d)

    for idx, d in enumerate(right_shapes_data):
        y = start_top + idx * (card_h + gap)
        render_shape_block(s, 6335000, y, 5303520, card_h, d)

    c_left = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 548640, 5440680, 5303520, 457200)
    c_left.fill.solid()
    c_left.fill.fore_color.rgb = COLOR_BG_CALLOUT
    c_left.line.color.rgb = COLOR_BORDER_BLUE
    c_left.line.width = Pt(1)
    pcl = c_left.text_frame.paragraphs[0]
    rcl = pcl.add_run()
    rcl.text = "💡 Golden Rule: Every flowchart must have exactly 1 Start and at least 1 End point!"
    rcl.font.name = FONT_CONTENT
    rcl.font.size = Pt(10.5)
    rcl.font.bold = True
    rcl.font.color.rgb = COLOR_BORDER_BLUE
    pcl.alignment = PP_ALIGN.CENTER

    c_right = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 6335000, 5440680, 5303520, 457200)
    c_right.fill.solid()
    c_right.fill.fore_color.rgb = COLOR_BG_CALLOUT
    c_right.line.color.rgb = COLOR_BORDER_AMBER
    c_right.line.width = Pt(1)
    pcr = c_right.text_frame.paragraphs[0]
    rcr = pcr.add_run()
    rcr.text = "⚡ Key Rule: Decision diamonds must always branch into at least 2 labeled paths!"
    rcr.font.name = FONT_CONTENT
    rcr.font.size = Pt(10.5)
    rcr.font.bold = True
    rcr.font.color.rgb = COLOR_YELLOW
    pcr.alignment = PP_ALIGN.CENTER

    return s

def build_presentation():
    if os.path.exists('Novonex Delta Slide.pptx'):
        prs = Presentation('Novonex Delta Slide.pptx')
        while len(prs.slides) > 0:
            rId = prs.slides._sldIdLst[0].rId
            prs.part.drop_rel(rId)
            del prs.slides._sldIdLst[0]
    else:
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

    print("Building polished presentation...")

    # ================= Slide 1: Title Slide =================
    create_title_slide(prs)

    # ================= Topic 1: Introduction to Programming & Compilers =================
    create_agenda_slide(
        prs, topic_num=1,
        topic_title="Introduction to Programming & Compilers",
        subtopics=["Programming", "Hardware", "Compilation", "Binaries", "Execution"]
    )
    create_content_slide(
        prs,
        badge_text="Introduction & Compilers",
        heading_text="Introduction to Programming & Compilers",
        bullets_data=[
            ("Programming", "Writing precise instructions that the CPU executes to solve problems."),
            ("Hardware", "Silicon CPU and physical memory (RAM) process binary electric states."),
            ("Compilation", "GCC translates human-readable high-level C code (.c) into machine code."),
            ("Binaries", "CPUs can only directly execute native machine binary instructions (0s and 1s)."),
            ("Execution", "OS loads compiled machine code into RAM for ultra-fast CPU execution.")
        ],
        callout_data=("💡", "Standard GCC Command: gcc -Wall main.c -o program", COLOR_BORDER_BLUE),
        right_board_data={
            "type": "diagram_box",
            "title": "✏️  Live Trace & Explanation Board",
            "subtitle": "[Teacher's Note: Architecture trace from high-level .c source to CPU execution]",
            "header": "From Source Code to Machine Execution:",
            "content": """1. Source File (.c):
   Human-readable C syntax statements and functions.

2. GCC Compilation Process:
   [ Preprocessor ] -> [ Compiler ] -> [ Assembler ] -> [ Linker ]

3. Executable Binary (.exe / ELF):
   Native 0s and 1s CPU instruction opcodes.

4. OS Loader & Physical RAM:
   Loads machine binary into memory; CPU executes at clock speed!

───────────────────────────────────────────────
• C bridges high-level logic and bare-metal hardware control."""
        }
    )
    create_empty_slide(prs)

    # ================= Topic 2: C Program Structure & Syntax (2 Slides) =================
    create_agenda_slide(
        prs, topic_num=2,
        topic_title="C Program Structure & Syntax",
        subtopics=["Headers", "Main", "Statements", "Blocks", "Comments"]
    )
    # Content Slide 1: Headers & Main Entry Point
    create_content_slide(
        prs,
        badge_text="Structure & Syntax",
        heading_text="C Program Structure & Execution Flow",
        bullets_data=[
            ("Headers", "Directives like #include <stdio.h> import standard I/O library prototypes."),
            ("Main", "Every C executable begins execution strictly at the mandatory int main()."),
            ("Statements", "Instructions in C must terminate with a semicolon (;) to mark completion.")
        ],
        code_snippet="""// Mandatory C Skeleton
#include <stdio.h>

int main() {
    printf("Hello, World!\\n");
    return 0; // Signals OS success
}""",
        callout_data=("💡", "Golden Rule: Execution always begins at main(), regardless of physical file position!", COLOR_BORDER_BLUE),
        right_board_data={
            "type": "diagram_box",
            "title": "✏️  Live Trace & Execution Flow",
            "subtitle": "[Line-by-line execution walkthrough of main entry point]",
            "header": "Program Execution Anatomy:",
            "content": """1. Preprocessor (#include):
   Pastes standard header declarations into compilation unit.

2. Function Header (int main()):
   Defines return type (int) and program entry point.

3. Standard Function Call (printf):
   Transfers bytes to stdout; '\\n' moves terminal cursor.

4. Return Statement (return 0):
   Returns status code 0 to operating system (Success!)."""
        }
    )
    # Content Slide 2: Blocks & Scoping Syntax
    create_content_slide(
        prs,
        badge_text="Structure & Syntax",
        heading_text="Blocks, Scopes & Documentation Syntax",
        bullets_data=[
            ("Blocks", "Curly braces { } define code blocks and isolate local variable scopes."),
            ("Comments", "Single-line // and multi-line /* */ document logic and are ignored by compilers."),
            ("Statements", "Compound statements grouped within braces execute as a single atomic unit.")
        ],
        code_snippet="""// Scope Block & Comments
int total = 100; // Global to block
{
    int local = 50; // Inner scope only
    total += local;
}
/* 'local' is destroyed after block! */""",
        callout_data=("💡", "Best Practice: Always use clear comments and consistent indentation for block scopes!", COLOR_BORDER_BLUE),
        right_board_data={
            "type": "dual_boxes",
            "title": "📦 Block Scope & Lifecycle Trace",
            "subtitle": "[How curly braces govern memory allocation and lifetime]",
            "box1_title": "1. Lexical Scope Isolation",
            "box1_content": """• Variables declared inside { } exist only in that block
• Outer scopes cannot access variables in inner blocks
• Prevents variable name collisions in large programs""",
            "box2_title": "2. Semicolon & Comment Rules",
            "box2_content": """• Semicolons terminate statements, NOT block braces { }
• Placing a semicolon right after 'if ();' creates an empty statement trap!
• Comments do not add any byte overhead to final binaries"""
        }
    )
    create_empty_slide(prs)

    # ================= Topic 3: Variables, Data Types & Standard I/O (2 Slides) =================
    create_agenda_slide(
        prs, topic_num=3,
        topic_title="Variables, Data Types & Standard I/O",
        subtopics=["Variables", "Types", "Specifiers", "Input", "Output"]
    )
    # Content Slide 1: Variables & Data Types
    create_content_slide(
        prs,
        badge_text="Variables & Data Types",
        heading_text="Variables, Data Types & Memory Allocation",
        bullets_data=[
            ("Variables", "Named memory locations in physical RAM allocated to store program data."),
            ("Types", "Fundamental types set size: int (4B), float (4B), double (8B), char (1B)."),
            ("Specifiers", "Format specifiers inform functions how to interpret raw binary bytes.")
        ],
        code_snippet="""// Data Type Declarations
int count = 42;         // 4 Bytes: Whole numbers
float price = 19.99f;   // 4 Bytes: 6 decimal places
double pi = 3.14159265; // 8 Bytes: 15 decimal places
char grade = 'A';       // 1 Byte:  ASCII character""",
        callout_data=("💡", "Memory Rule: sizeof(type) queries the exact byte footprint on the target hardware!", COLOR_BORDER_BLUE),
        right_board_data={
            "type": "diagram_box",
            "title": "🧠 RAM Memory Box Architecture",
            "subtitle": "[Hardware memory allocation and bit representations]",
            "header": "Physical Memory Architecture:",
            "content": """Variable Storage in Hardware RAM:
┌───────────┬──────────┬───────────┬──────────────────┐
│ Variable  │ Type     │ Size (B)  │ Memory Address   │
├───────────┼──────────┼───────────┼──────────────────┤
│ count     │ int      │ 4 Bytes   │ 0x7ffd5e20 (RAM) │
│ price     │ float    │ 4 Bytes   │ 0x7ffd5e24 (RAM) │
│ pi        │ double   │ 8 Bytes   │ 0x7ffd5e28 (RAM) │
│ grade     │ char     │ 1 Byte    │ 0x7ffd5e30 (RAM) │
└───────────┴──────────┴───────────┴──────────────────┘

• Variables map directly to contiguous memory cells
• Hexadecimal addresses locate each variable in RAM"""
        }
    )
    # Content Slide 2: Standard I/O
    create_content_slide(
        prs,
        badge_text="Standard I/O",
        heading_text="Standard Console Input & Output Streams",
        bullets_data=[
            ("Specifiers", "Tokens (%d, %f, %lf, %c) direct formatted printing and reading conversions."),
            ("Input", "scanf() captures user terminal keystrokes and requires the memory address (&)."),
            ("Output", "printf() writes formatted character streams directly to standard console.")
        ],
        code_snippet="""// Formatted Console Reading & Writing
int age; double gpa;
printf("Enter age and gpa: ");
scanf("%d %lf", &age, &gpa);
printf("Age = %d | GPA = %.2lf\\n", age, gpa);""",
        callout_data=("⚠️", "Critical Trap: Forgetting '&' in scanf() causes segmentation fault crashes!", COLOR_BORDER_AMBER),
        right_board_data={
            "type": "dual_boxes",
            "title": "⌨️  Standard I/O Buffer Tracing",
            "subtitle": "[How printf and scanf interact with stdin/stdout]",
            "box1_title": "1. Why scanf Requires '&'",
            "box1_content": """• C passes arguments by value (copies)
• To modify 'age', scanf must receive its address (&age)
• Without '&', scanf interprets age's value as an address!""",
            "box2_title": "2. Common Format Specifiers",
            "box2_content": """• %d  -> Signed decimal integer (int)
• %f  -> Single-precision float
• %lf -> Double-precision float (in scanf)
• %c  -> Single character (char)"""
        }
    )
    create_empty_slide(prs)

    # ================= Topic 4: Operators, Identifiers, & Expression Evaluation (2 Slides) =================
    create_agenda_slide(
        prs, topic_num=4,
        topic_title="Operators, Identifiers, & Expression Evaluation",
        subtopics=["Identifiers", "Arithmetic", "Relational", "Logical", "Precedence"]
    )
    # Content Slide 1: Identifiers & Basic Operators
    create_content_slide(
        prs,
        badge_text="Operators & Identifiers",
        heading_text="Identifiers & Basic Operator Mechanics",
        bullets_data=[
            ("Identifiers", "Programmer-defined variable and function names following strict naming rules."),
            ("Arithmetic", "Math operators: +, -, *, / (quotient), and % (modulo integer remainder)."),
            ("Relational", "Comparison operators (==, !=, <, >, <=, >=) evaluate to 1 (True) or 0 (False).")
        ],
        code_snippet="""// Arithmetic & Relational Evaluation
int a = 7, b = 2;
int quot = a / b;     // 3 (integer truncation!)
int rem = a % b;      // 1 (remainder)
int is_greater = a > b; // 1 (True)""",
        callout_data=("⚠️", "Integer Division Trap: 7 / 2 yields 3! Use 7.0 / 2 to obtain 3.5 float result!", COLOR_BORDER_AMBER),
        right_board_data={
            "type": "diagram_box",
            "title": "✏️  Live Operator Evaluation Board",
            "subtitle": "[Identifier rules & arithmetic truncation behavior]",
            "header": "Operator Rules & Identifiers:",
            "content": """1. Valid Identifier Rules:
   • Must start with letter (a-z, A-Z) or underscore (_)
   • Cannot start with a digit (e.g. 2var is INVALID)
   • Case-sensitive: 'Total' and 'total' are distinct!

2. Modulo Operator (%) Constraints:
   • Only works with integer operands (not float/double)
   • Useful for parity checks: (n % 2 == 0) -> Even!

3. Relational Truth Values:
   • In C, comparisons yield int 1 for TRUE and 0 for FALSE."""
        }
    )
    # Content Slide 2: Logical Operators & Precedence
    create_content_slide(
        prs,
        badge_text="Operators & Precedence",
        heading_text="Logical Operators & Precedence Ladder",
        bullets_data=[
            ("Logical", "Boolean operators (&&, ||, !) enable compound condition testing with short-circuiting."),
            ("Precedence", "Precedence ranking resolves operator priority in compound mathematical expressions."),
            ("Relational", "Relational evaluations bind tighter than logical tests but looser than arithmetic.")
        ],
        code_snippet="""// Compound Logic & Precedence
int x = 5, y = 10;
// && has higher precedence than ||
int res = (x > 0) && (y < 20); // 1 (True)
int calc = x + y * 2;          // 25 (* precedes +)""",
        callout_data=("💡", "Golden Rule: Always use parentheses () to make operator precedence explicit!", COLOR_BORDER_BLUE),
        right_board_data={
            "type": "ladder_boxes",
            "title": "📊 Graphical Operator Precedence Ladder",
            "subtitle": "[Priority hierarchy & evaluation direction]",
            "border_color": COLOR_BORDER_BLUE,
            "items": [
                "1. Parentheses: ()  []  ->  .  | Left to Right",
                "2. Unary Operators: ++  --  !  sizeof  &  | Right to Left",
                "3. Multiplicative: *  /  %  | Left to Right",
                "4. Additive: +  -  | Left to Right",
                "5. Relational: <  <=  >  >=  | Left to Right",
                "6. Equality: ==  !=  | Left to Right",
                "7. Logical AND / OR: && then ||  | Left to Right",
                "8. Assignment: =  +=  -=  *=  | Right to Left"
            ]
        }
    )
    create_empty_slide(prs)

    # ================= Topic 5: Conditional Statements & Nested Logic (2 Slides) =================
    create_agenda_slide(
        prs, topic_num=5,
        topic_title="Conditional Statements (if, else if, else) and Nested Logic",
        subtopics=["Branching", "Conditions", "Alternatives", "Nesting", "Fallthrough"]
    )
    # Content Slide 1: if, else if, else branching
    create_content_slide(
        prs,
        badge_text="Conditionals: if-else",
        heading_text="Conditional Statements: if, else if, else",
        bullets_data=[
            ("Branching", "Alters sequential instruction flow based on runtime evaluation of conditions."),
            ("Conditions", "Expressions where 0 evaluates to False and any non-zero value evaluates to True."),
            ("Alternatives", "else if chains and trailing else blocks guarantee mutual exclusion across paths.")
        ],
        code_snippet="""// Multi-Way Branching: Grade Classifier
if (score >= 90) {
    printf("Grade: A\\n");
} else if (score >= 80) {
    printf("Grade: B\\n");
} else {
    printf("Grade: C\\n");
}""",
        callout_data=("⚠️", "Danger: if (x = 5) assigns 5 (always TRUE!). Always use '==' for comparison checks!", COLOR_BORDER_AMBER),
        right_board_data={
            "type": "diagram_box",
            "title": "✏️  Live Trace & Decision Trees",
            "subtitle": "[Trace decision tree: Single if vs if-else vs else-if ladder]",
            "header": "Decision Logic Architecture:",
            "content": """1. Single if:
   [ Condition True? ] ──Yes──> [ Execute Body ]
           │ No
           ▼ (Continue sequential flow)

2. Two-way if-else:
   [ Condition True? ] ──Yes──> [ True Branch ]
           │ No
           └─────────> [ False Branch ]

3. Multi-way else-if Ladder:
   Tests conditions sequentially from top to bottom;
   first matching True block executes, rest skipped!"""
        }
    )
    # Content Slide 2: Nested Logic & Switch Case
    create_content_slide(
        prs,
        badge_text="Nested Logic & Switch",
        heading_text="Nested Conditionals & Switch Logic",
        bullets_data=[
            ("Nesting", "Embedding an if-else inside another conditional tests granular multi-tier logic."),
            ("Fallthrough", "In switch blocks, execution cascades into next cases unless halted with break;."),
            ("Alternatives", "The default: branch reliably captures any unexpected or unmatched inputs.")
        ],
        code_snippet="""// Multi-Way Decision: switch with break
switch (op) {
    case '+': res = a + b; break;
    case '-': res = a - b; break;
    case '*': res = a * b; break;
    default:  printf("Invalid operator\\n");
}""",
        callout_data=("💡", "Best Practice: Flatten deep nested conditionals using guard clauses and early exits!", COLOR_BORDER_BLUE),
        right_board_data={
            "type": "dual_boxes",
            "title": "⚡ Decision Tree vs Switch Jump Table",
            "subtitle": "[Architectural comparison of multi-branching]",
            "box1_title": "1. Nested If-Else Decision Tree",
            "box1_content": """• Checks conditions sequentially: O(N) comparisons
• Ideal for ranges (score >= 80 && score < 90)
• High nesting levels reduce code readability""",
            "box2_title": "2. Switch-Case Jump Table",
            "box2_content": """• Compiler compiles into direct O(1) jump table
• Strictly tests discrete integer or char values
• Missing 'break;' triggers case fallthrough!"""
        }
    )
    create_empty_slide(prs)

    # ================= Topic 6: Flow chart (2 Slides) =================
    create_agenda_slide(
        prs, topic_num=6,
        topic_title="Flow chart",
        subtopics=["Flowcharts", "Symbols", "Processes", "Decisions", "Patterns"]
    )
    # Content Slide 1: Flowchart Shapes & Standard Symbols
    create_flowchart_shapes_slide(prs)

    # Content Slide 2: Flowchart Fundamentals & Logic Flow
    create_content_slide(
        prs,
        badge_text="Flow Charts",
        heading_text="Flowchart Fundamentals & Logic Flow",
        bullets_data=[
            ("Flowcharts", "Standardized visual blueprints representing algorithmic step-by-step program logic."),
            ("Purpose", "Clarifies complex logic, prevents syntax & logic bugs, and acts as an architectural guide."),
            ("Execution", "Control flows strictly along directed lines from the Start node to an End node."),
            ("Dry Running", "Enables step-by-step manual tracing of variables with test data prior to coding."),
            ("Golden Rules", "Must possess unambiguous pathways, a single entry terminal, and reachable exits.")
        ],
        callout_data=("💡", "Best Practice: Design and trace your flowchart logic on paper before writing actual C code!", COLOR_BORDER_BLUE),
        right_board_data={
            "type": "diagram_box",
            "title": "📐 Core Flowchart Control Patterns",
            "subtitle": "[Visual representation of 3 fundamental algorithmic control structures]",
            "header": "ANSI Flowchart Patterns:",
            "content": """1. Sequence Pattern (Linear Progression):
   [ Start ] ──> [ Input A, B ] ──> [ Sum = A + B ] ──> [ End ]

2. Selection Pattern (if-else branching):
                  [ Condition? ]
                  /            \\
           (True)/              \\(False)
                ▼                ▼
         [ Body Block ]    [ Else Block ]
                \\                /
                 ───────>•<───────
                         ▼
3. Loop Pattern (Iteration cycle):
   ──> [ Check Condition ] ──False──> [ Exit Loop ]
             │ True
             ▼
       [ Loop Body ] ──> [ Update ] ──┘"""
        }
    )
    create_empty_slide(prs)

    # ================= Topic 7: Iterative Statements & Nested Loops (2 Slides) =================
    create_agenda_slide(
        prs, topic_num=7,
        topic_title="Iterative Statements (Loops) and Nested Loops",
        subtopics=["Iteration", "Initialization", "Condition", "Update", "Nesting"]
    )
    # Content Slide 1: Single Loops
    create_content_slide(
        prs,
        badge_text="Iterative Loops",
        heading_text="Iterative Statements: for, while & do-while",
        bullets_data=[
            ("Iteration", "Repeated execution of an instruction block until a stopping condition is met."),
            ("Initialization", "Counter variables are initialized once prior to entering the loop cycle."),
            ("Condition", "Boolean expression verified before each pass; loop halts immediately when False."),
            ("Update", "Loop counter is modified (e.g. i++) after every iteration to progress toward exit.")
        ],
        code_snippet="""// Counter Controlled Loop (for)
for (int i = 0; i < 5; i++) {
    printf("%d ", i);  // Prints: 0 1 2 3 4
}
// Condition Controlled Loop (while)
int n = 10;
while (n > 0) { n /= 2; }""",
        callout_data=("💡", "Execution Order: 1. Init -> 2. Condition Check -> 3. Body Execute -> 4. Counter Update", COLOR_BORDER_BLUE),
        right_board_data={
            "type": "dual_boxes",
            "title": "🔄 while vs do-while Flow Comparison",
            "subtitle": "[Entry-controlled vs Exit-controlled architecture]",
            "box1_title": "1. while Loop (Entry-Controlled Flow)",
            "box1_content": """   [ Check Condition First ]
             |── True ──> [ Execute Body ] ──┘
             └── False ─> [ Exit Immediately ]
• May execute 0 times if condition starts FALSE!""",
            "box2_title": "2. do-while Loop (Exit-Controlled Flow)",
            "box2_content": """   [ Execute Body FIRST without checking ]
             |
             v
   [ Check Condition at Bottom ]
             |── True ──> [ Repeat Body ]
             └── False ─> [ Exit Loop ]
• Guaranteed to execute at least 1 time!"""
        }
    )
    # Content Slide 2: Nested Loops
    create_content_slide(
        prs,
        badge_text="Nested Loops",
        heading_text="Nested Loops & Grid Traversal",
        bullets_data=[
            ("Nesting", "A loop placed entirely inside another loop creates multi-dimensional iteration cycles."),
            ("Iteration", "The inner loop runs completely from start to finish for every iteration of outer loop."),
            ("Update", "Outer loop updates row indices, while inner loop steps through column coordinates.")
        ],
        code_snippet="""// Pattern Example: 3 Rows x 4 Stars Grid
for (int r = 1; r <= 3; r++) {     // Outer controls Rows
    for (int c = 1; c <= 4; c++) { // Inner controls Columns
        printf("* ");
    }
    printf("\\n");
}""",
        callout_data=("💡", "Golden Rule: Outer loop always controls Rows (r); Inner loop controls Columns (c)!", COLOR_BORDER_BLUE),
        right_board_data={
            "type": "diagram_box",
            "title": "📐 Nested Loop Matrix Grid Traversal",
            "subtitle": "[How 2 loops generate 2D grids and patterns]",
            "header": "Visual Execution Trace: 3 Rows x 4 Columns",
            "content": """Outer (r=1)  -> Inner (c=1..4):   *  *  *  *   [Newline]
Outer (r=2)  -> Inner (c=1..4):   *  *  *  *   [Newline]
Outer (r=3)  -> Inner (c=1..4):   *  *  *  *   [Newline]

───────────────────────────────────────────────
Execution Step Trace:
• Outer loop runs 3 times (r = 1, 2, 3)
• For each r, inner loop runs 4 times (c = 1, 2, 3, 4)
• Total printf("* ") calls = 3 x 4 = 12 times!
• Total time complexity = O(Rows x Cols)"""
        }
    )
    create_empty_slide(prs)

    # ================= Topic 8: 1D Arrays: Memory Layout & Indexing (2 Slides) =================
    create_agenda_slide(
        prs, topic_num=8,
        topic_title="1D Arrays: Memory Layout & Indexing",
        subtopics=["Arrays", "Indexing", "Contiguity", "Addresses", "Bounds"]
    )
    # Content Slide 1: Fundamentals & Indexing
    create_content_slide(
        prs,
        badge_text="1D Arrays",
        heading_text="1D Arrays: Definition & Zero-Based Indexing",
        bullets_data=[
            ("Arrays", "Fixed-size collections storing elements of identical data type in contiguous memory."),
            ("Indexing", "Zero-based indexing (0 to N-1) grants direct, instant O(1) random access to elements."),
            ("Bounds", "C performs no runtime bounds checks; accessing beyond bounds leads to memory corruption.")
        ],
        code_snippet="""// 1D Array Declaration, Input & Output
int arr[5] = {10, 20, 30, 40, 50};
for (int i = 0; i < 5; i++) {
    printf("arr[%d] = %d\\n", i, arr[i]);
}""",
        callout_data=("⚠️", "Out of Bounds: Accessing arr[5] in a size-5 array leads to undefined behavior / garbage!", COLOR_BORDER_AMBER),
        right_board_data={
            "type": "diagram_box",
            "title": "🧠 1D Array Contiguous RAM Layout",
            "subtitle": "[Physical memory allocation in hardware]",
            "header": "Visual RAM Memory Cells: int arr[5] = {10, 20, 30, 40, 50}",
            "content": """  Index:       [0]      [1]      [2]      [3]      [4]
             ┌──────┬──────┬──────┬──────┬──────┐
  Value:     │  10  │  20  │  30  │  40  │  50  │
             └──────┴──────┴──────┴──────┴──────┘
  Address:    1000   1004   1008   1012   1016   (+4B per int)

───────────────────────────────────────────────
• Base Address: Address of arr[0] = 1000 (&arr[0] or arr)
• Formula: Address(arr[i]) = Base + (i x sizeof(type))
  Example: Address(arr[3]) = 1000 + (3 x 4) = 1012
• Contiguous: Zero gaps between elements in RAM!"""
        }
    )
    # Content Slide 2: Memory Contiguity & Address Calculation
    create_content_slide(
        prs,
        badge_text="1D Arrays",
        heading_text="1D Arrays: Memory Layout & Pointer Offsets",
        bullets_data=[
            ("Contiguity", "Array elements occupy adjacent, unbroken bytes in physical RAM without any gaps."),
            ("Addresses", "Memory location computed via hardware formula: Address(arr[i]) = Base + (i * sizeof(type))."),
            ("Bounds", "Array name acts as a constant pointer to the base address of the first element (&arr[0]).")
        ],
        code_snippet="""// Pointer Arithmetic & Array Addressing
int arr[3] = {100, 200, 300};
printf("Base Address: %p\\n", (void*)arr);
printf("arr[1] Addr:  %p\\n", (void*)&arr[1]);
// Address increases by sizeof(int) = 4 bytes!""",
        callout_data=("💡", "Golden Rule: Array indexing arr[i] is syntactically equivalent to *(arr + i)!", COLOR_BORDER_BLUE),
        right_board_data={
            "type": "dual_boxes",
            "title": "📐 Hardware Pointer Offset Mechanics",
            "subtitle": "[How the CPU computes array element memory addresses]",
            "box1_title": "1. Linear Address Calculation",
            "box1_content": """• Base Address = starting byte of arr[0] in RAM
• Offset formula = Base + (index * byte_size)
• Multiplication is done in hardware using bit shifts
• Gives constant O(1) random access time!""",
            "box2_title": "2. The Danger of Index Out-of-Bounds",
            "box2_content": """• arr[-1] or arr[5] reads adjacent stack memory
• Can overwrite loop counters, return addresses
• Major source of critical security exploits!"""
        }
    )
    create_empty_slide(prs)

    # ================= Topic 9: 2D Arrays (Matrices): Grid & Memory Mapping (2 Slides) =================
    create_agenda_slide(
        prs, topic_num=9,
        topic_title="2D Arrays (Matrices): Grid & Memory Mapping",
        subtopics=["Matrices", "Rows", "Columns", "Traversal", "Linearization"]
    )
    # Content Slide 1: Grid Concepts & Traversal
    create_content_slide(
        prs,
        badge_text="2D Arrays (Matrices)",
        heading_text="2D Arrays: Grid Coordinates & Matrix Model",
        bullets_data=[
            ("Matrices", "Two-dimensional arrays structured conceptually as a table grid of Rows and Columns."),
            ("Rows", "The primary dimension index referencing the horizontal row partition of matrix data."),
            ("Columns", "The secondary dimension index accessing individual element cells along horizontal rows.")
        ],
        code_snippet="""// 2D Array Declaration & Grid Setup
int mat[2][3] = {
    {1, 2, 3},  // Row 0
    {4, 5, 6}   // Row 1
};
printf("Cell (0,1): %d\\n", mat[0][1]); // 2""",
        callout_data=("💡", "Convention: In mat[R][C], first index R specifies Row, second index C specifies Column!", COLOR_BORDER_BLUE),
        right_board_data={
            "type": "diagram_box",
            "title": "📐 2D Coordinate Grid Matrix",
            "subtitle": "[Logical table visualization of rows and columns]",
            "header": "Conceptual 2D Grid (2 Rows x 3 Cols):",
            "content": """             Col 0     Col 1     Col 2
   Row 0:   [ (0,0) ] [ (0,1) ] [ (0,2) ]
   Row 1:   [ (1,0) ] [ (1,1) ] [ (1,2) ]

───────────────────────────────────────────────
Logical Cell Values:
• mat[0][0] = 1,  mat[0][1] = 2,  mat[0][2] = 3
• mat[1][0] = 4,  mat[1][1] = 5,  mat[1][2] = 6
• Total cells = ROWS x COLS = 2 x 3 = 6 integers!"""
        }
    )
    # Content Slide 2: Row-Major Memory Mapping & Traversal
    create_content_slide(
        prs,
        badge_text="2D Arrays (Matrices)",
        heading_text="2D Arrays: Traversal & Row-Major RAM Mapping",
        bullets_data=[
            ("Traversal", "Nested loops navigate 2D arrays: outer loop steps rows (i), inner loop steps columns (j)."),
            ("Linearization", "Physical RAM is linear; C flattens 2D matrices into 1D memory using Row-Major order."),
            ("Rows", "Row 0 elements are stored consecutively in RAM, followed immediately by Row 1 elements.")
        ],
        code_snippet="""// 2D Array Traversal & Address Tracing
for (int i = 0; i < 2; i++) {
    for (int j = 0; j < 3; j++) {
        printf("mat[%d][%d]=%d ", i, j, mat[i][j]);
    }
    printf("\\n");
}""",
        callout_data=("💡", "Linear Offset Formula: Address(mat[i][j]) = Base + (i x COLS + j) x sizeof(type)", COLOR_BORDER_BLUE),
        right_board_data={
            "type": "diagram_box",
            "title": "🧠 2D Grid to Row-Major RAM Mapping",
            "subtitle": "[Hardware linearization of 2D matrices]",
            "header": "Row-Major Memory Order in Physical RAM:",
            "content": """1. Conceptual 2D Grid (2 Rows x 3 Cols):
             Col 0   Col 1   Col 2
   Row 0:   [  1  ] [  2  ] [  3  ]
   Row 1:   [  4  ] [  5  ] [  6  ]

2. Physical Linear RAM Storage (Row-Major):
   [ 1 ] [ 2 ] [ 3 ]   -->   [ 4 ] [ 5 ] [ 6 ]
   └───── Row 0 ────┘         └───── Row 1 ────┘

───────────────────────────────────────────────
• Row 0 is stored consecutively, followed immediately by Row 1
• Row-wise iteration maximizes CPU cache hit rates!"""
        }
    )
    create_empty_slide(prs)

    # ================= Topic 10: Error finding in C code (2 Slides) =================
    create_agenda_slide(
        prs, topic_num=10,
        topic_title="Error finding in C code",
        subtopics=["Syntax", "Linking", "Runtime", "Logic", "Debugging"]
    )
    # Content Slide 1: Syntax & I/O Traps
    create_content_slide(
        prs,
        badge_text="Error Finding Lab",
        heading_text="Spot the Bug: Syntax & Input/Output Traps",
        bullets_data=[
            ("Syntax", "Grammatical mistakes violating C rules like missing semicolons or case errors (Printf)."),
            ("Linking", "Unresolved symbol references occurring when function declarations have no definitions."),
            ("Debugging", "Enabling aggressive compiler warnings (-Wall -Wextra) exposes syntax bugs instantly.")
        ],
        code_snippet="""// Snippet A: Reading integer from user
int age; scanf("%d", age);           // BUG #1: Missing &
// Snippet B: Score output
int score = 95
Printf("Score: %d\\n", score);       // BUGS #2 & #3: Semicolon & Printf
// Snippet C: Float printing
float gpa = 3.85; printf("%d", gpa); // BUG #4: %d with float""",
        callout_data=("🔍", "Student Challenge: Identify the 4 errors above and explain the compiler errors!", COLOR_BORDER_AMBER),
        right_board_data={
            "type": "diagram_box",
            "title": "✏️  Live Trace & Bug Fix Board",
            "subtitle": "[Classroom Activity: Ask students to spot errors, circle them live, and write corrected code]",
            "header": "Bug Diagnosis & Fix Guide:",
            "content": """Bug #1 Fix:
• scanf("%d", &age); // Pass memory address with &!

Bug #2 & #3 Fix:
• int score = 95;    // Semicolon required after statement!
• printf("Score..."); // C is case-sensitive: 'printf', not 'Printf'!

Bug #4 Fix:
• printf("GPA: %.2f\\n", gpa); // Use %f (or %lf) for floating point!

───────────────────────────────────────────────
• Always compile with 'gcc -Wall -Wextra' to catch these early!"""
        }
    )
    # Content Slide 2: Logic, Loops & Array Traps
    create_content_slide(
        prs,
        badge_text="Error Finding Lab",
        heading_text="Spot the Bug: Logic, Loops & Array Traps",
        bullets_data=[
            ("Logic", "Using assignment '=' instead of equality '==' in conditionals always evaluates to True."),
            ("Runtime", "Off-by-one array boundary violations (i <= N) read or overwrite garbage memory."),
            ("Debugging", "Tracing variable states across loop boundary iterations quickly exposes off-by-one flaws.")
        ],
        code_snippet="""// Snippet D: Conditional Security Check
int is_admin = 0;
if (is_admin = 1) {       // BUG #5: Assignment '=' vs '=='
    printf("Admin Access Granted!\\n");
}
// Snippet E: Array Traversal & Printing
int arr[5] = {10, 20, 30, 40, 50};
for (int i = 0; i <= 5; ) { // BUGS #6 & #7: Bounds <= 5 & missing i++
    printf("%d ", arr[i]);
}""",
        callout_data=("🔍", "Student Challenge: These compile without error but give wrong results! Why?", COLOR_BORDER_AMBER),
        right_board_data={
            "type": "diagram_box",
            "title": "✏️  Live Trace & Bug Fix Board",
            "subtitle": "[Classroom Activity: Trace condition truth table, index bounds, and float cast]",
            "header": "Subtle Logic Bug Fixes:",
            "content": """Bug #5 Fix (Assignment vs Equality):
• if (is_admin == 1) // '=' assigns 1 (always TRUE); use '=='!

Bug #6 & #7 Fix (Array Bounds & Infinite Loop):
• for (int i = 0; i < 5; i++) // '< 5' prevents arr[5] crash, 'i++' updates!

General Debugging Practice:
• Trace loops using a dry-run iteration table
• Check boundary edges (index 0 and index N-1)
• Treat compiler warnings as fatal errors (-Werror)!"""
        }
    )
    create_empty_slide(prs)

    # ================= Final Slide: Thank You Slide =================
    create_thank_you_slide(prs)

    # Save presentation
    prs.save(OUTPUT_FILE)
    print(f"Successfully generated '{OUTPUT_FILE}' with {len(prs.slides)} slides!")

if __name__ == '__main__':
    build_presentation()
