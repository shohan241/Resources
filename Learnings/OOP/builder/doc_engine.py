import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

class DocEngine:
    def __init__(self, docx_path, md_path):
        self.docx_path = docx_path
        self.md_path = md_path
        self.doc = docx.Document()
        self.md_lines = []
        
        # Color definitions for executive layout
        self.C_NAVY = RGBColor(15, 39, 68)       # #0F2744
        self.C_BLUE = RGBColor(30, 90, 168)      # #1E5AA8
        self.C_TEXT = RGBColor(30, 41, 59)       # #1E293B
        self.C_MUTED = RGBColor(100, 116, 139)   # #64748B
        self.C_CODE_TXT = RGBColor(241, 245, 249)# #F1F5F9
        self.C_CODE_BG = "0F172A"                # Dark Navy Slate
        self.C_BORDER = "CBD5E1"
        self.C_BG_CARD = "F1F5F9"
        self.C_BG_ALERT = "FEF2F2"
        self.C_BG_SUCCESS = "F0FDF4"
        self.C_BG_WARNING = "FFFBEB"

        self._setup_page()

    def _setup_page(self):
        for sec in self.doc.sections:
            sec.top_margin = Inches(0.7)
            sec.bottom_margin = Inches(0.7)
            sec.left_margin = Inches(0.75)
            sec.right_margin = Inches(0.75)
            
            # Header
            hp = sec.header.paragraphs[0]
            hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            hr = hp.add_run("Java OOP & Low-Level Design (LLD) | Complete Industry Mastery Guide")
            hr.font.name = "Calibri"
            hr.font.size = Pt(8.5)
            hr.font.color.rgb = RGBColor(148, 163, 184)
            
            # Footer
            fp = sec.footer.paragraphs[0]
            fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fr = fp.add_run("Java Object-Oriented Software Engineering & Systems Design Handbook")
            fr.font.name = "Calibri"
            fr.font.size = Pt(8.5)
            fr.font.color.rgb = RGBColor(148, 163, 184)

    def add_title(self, title, subtitle):
        # DOCX Title
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(title)
        r.font.name = "Calibri"
        r.font.size = Pt(22)
        r.font.bold = True
        r.font.color.rgb = self.C_NAVY

        # Subtitle
        p_sub = self.doc.add_paragraph()
        p_sub.paragraph_format.space_before = Pt(0)
        p_sub.paragraph_format.space_after = Pt(14)
        r_sub = p_sub.add_run(subtitle)
        r_sub.font.name = "Calibri"
        r_sub.font.size = Pt(10)
        r_sub.font.italic = True
        r_sub.font.color.rgb = self.C_MUTED

        # Markdown
        self.md_lines.append(f"# {title}\n")
        self.md_lines.append(f"*{subtitle}*\n\n---\n")

    def add_chapter(self, num, title, overview=""):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        
        rn = p.add_run(f"Chapter {num}: ")
        rn.font.name = "Calibri"
        rn.font.size = Pt(14)
        rn.font.bold = True
        rn.font.color.rgb = self.C_BLUE
        
        rt = p.add_run(title)
        rt.font.name = "Calibri"
        rt.font.size = Pt(14)
        rt.font.bold = True
        rt.font.color.rgb = self.C_NAVY

        # Markdown
        self.md_lines.append(f"\n## Chapter {num}: {title}\n")

        if overview:
            self.add_body(overview, space_after=6)

    def add_h2(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(11.5)
        r.font.bold = True
        r.font.color.rgb = self.C_BLUE

        self.md_lines.append(f"\n### {text}\n")

    def add_h3(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = self.C_TEXT

        self.md_lines.append(f"\n#### {text}\n")

    def add_body(self, text, bold_prefix=None, space_after=3):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            prefix = bold_prefix if bold_prefix.endswith(" ") else bold_prefix + " "
            rp = p.add_run(prefix)
            rp.font.name = "Calibri"
            rp.font.size = Pt(9.5)
            rp.font.bold = True
            rp.font.color.rgb = self.C_TEXT
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r.font.color.rgb = self.C_TEXT

        # Markdown
        prefix_md = f"**{bold_prefix}** " if bold_prefix else ""
        self.md_lines.append(f"{prefix_md}{text}\n")

    def add_bullet(self, bold_prefix, text):
        p = self.doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        
        prefix = bold_prefix if bold_prefix.endswith(" ") else bold_prefix + " "
        rp = p.add_run(prefix)
        rp.font.name = "Calibri"
        rp.font.size = Pt(9)
        rp.font.bold = True
        rp.font.color.rgb = self.C_TEXT
        
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(9)
        r.font.color.rgb = self.C_TEXT

        self.md_lines.append(f"- **{bold_prefix}** {text}")

    def add_code(self, title, lang, code_str):
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.width = Inches(7.0)
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{self.C_CODE_BG}"/>')
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="6" w:space="0" w:color="334155"/><w:left w:val="single" w:sz="18" w:space="0" w:color="00AEEF"/><w:bottom w:val="single" w:sz="6" w:space="0" w:color="334155"/><w:right w:val="single" w:sz="6" w:space="0" w:color="334155"/></w:tcBorders>')
        tcPr.append(shd)
        tcPr.append(borders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(1)
        if title:
            rt = p.add_run(f"// --- {title} ({lang.upper()}) ---\n")
            rt.font.name = "Consolas"
            rt.font.size = Pt(8.5)
            rt.font.bold = True
            rt.font.color.rgb = RGBColor(56, 189, 248)
            
        rc = p.add_run(code_str.strip())
        rc.font.name = "Consolas"
        rc.font.size = Pt(8)
        rc.font.color.rgb = self.C_CODE_TXT
        
        p_sp = self.doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(0)
        p_sp.paragraph_format.space_after = Pt(3)

        # Markdown
        self.md_lines.append(f"\n```{lang}\n// {title}\n{code_str.strip()}\n```\n")

    def add_callout(self, title, lines, box_type="info"):
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.width = Inches(7.0)
        tcPr = cell._tc.get_or_add_tcPr()
        
        if box_type == "alert":
            bg = self.C_BG_ALERT
            bar = "DC2626"
            t_col = RGBColor(185, 28, 28)
            md_tag = "CAUTION"
        elif box_type == "success":
            bg = self.C_BG_SUCCESS
            bar = "16A34A"
            t_col = RGBColor(22, 101, 52)
            md_tag = "TIP"
        elif box_type == "warning":
            bg = self.C_BG_WARNING
            bar = "D97706"
            t_col = RGBColor(180, 83, 9)
            md_tag = "WARNING"
        else:
            bg = self.C_BG_CARD
            bar = "1E5AA8"
            t_col = self.C_BLUE
            md_tag = "NOTE"
            
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg}"/>')
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="{self.C_BORDER}"/><w:left w:val="single" w:sz="24" w:space="0" w:color="{bar}"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="{self.C_BORDER}"/><w:right w:val="single" w:sz="4" w:space="0" w:color="{self.C_BORDER}"/></w:tcBorders>')
        tcPr.append(shd)
        tcPr.append(borders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(1.5)
        rt = p.add_run(f"★ {title}")
        rt.font.name = "Calibri"
        rt.font.size = Pt(9.5)
        rt.font.bold = True
        rt.font.color.rgb = t_col
        
        for l in lines:
            pl = cell.add_paragraph()
            pl.paragraph_format.space_before = Pt(0)
            pl.paragraph_format.space_after = Pt(1.5)
            pl.paragraph_format.line_spacing = 1.15
            rl = pl.add_run(l)
            rl.font.name = "Calibri"
            rl.font.size = Pt(8.5)
            rl.font.color.rgb = self.C_TEXT
            
        p_sp = self.doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(0)
        p_sp.paragraph_format.space_after = Pt(3)

        # Markdown
        self.md_lines.append(f"\n> [!{md_tag}]\n> **{title}**\n" + "".join([f"> {l}\n" for l in lines]))

    def add_table(self, headers, rows, widths=None):
        if widths is None:
            total_w = 7.0
            col_w = Inches(total_w / len(headers))
            widths = [col_w] * len(headers)

        table = self.doc.add_table(rows=len(rows) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_cells = table.rows[0].cells
        
        for i, h in enumerate(headers):
            hdr_cells[i].width = widths[i]
            p = hdr_cells[i].paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(h)
            run.font.name = "Calibri"
            run.font.size = Pt(8.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            
            tcPr = hdr_cells[i]._tc.get_or_add_tcPr()
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="0F2744"/>')
            borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="0" w:color="0A192F"/><w:top w:val="none"/><w:left w:val="none"/><w:right w:val="none"/></w:tcBorders>')
            tcPr.append(shd)
            tcPr.append(borders)
            
        for r_idx, row in enumerate(rows):
            row_cells = table.rows[r_idx + 1].cells
            bg_col = "FFFFFF" if r_idx % 2 == 0 else "F8FAFC"
            for c_idx, val in enumerate(row):
                row_cells[c_idx].width = widths[c_idx]
                p = row_cells[c_idx].paragraphs[0]
                p.paragraph_format.space_before = Pt(2.5)
                p.paragraph_format.space_after = Pt(2.5)
                p.paragraph_format.line_spacing = 1.05
                run = p.add_run(str(val))
                run.font.name = "Calibri"
                run.font.size = Pt(8)
                run.font.color.rgb = self.C_TEXT
                if c_idx == 0:
                    run.font.bold = True
                tcPr = row_cells[c_idx]._tc.get_or_add_tcPr()
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_col}"/>')
                borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:space="0" w:color="{self.C_BORDER}"/><w:top w:val="none"/><w:left w:val="none"/><w:right w:val="none"/></w:tcBorders>')
                tcPr.append(shd)
                tcPr.append(borders)
                
        p_sp = self.doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(0)
        p_sp.paragraph_format.space_after = Pt(4)

        # Markdown Table
        md_hdr = "| " + " | ".join(headers) + " |"
        md_sep = "| " + " | ".join(["---"] * len(headers)) + " |"
        md_rows = ["| " + " | ".join([str(c).replace("\n", "<br>") for c in row]) + " |" for row in rows]
        self.md_lines.append("\n" + md_hdr + "\n" + md_sep + "\n" + "\n".join(md_rows) + "\n")

    def save(self):
        print(f"Saving DOCX to {self.docx_path}...")
        self.doc.save(self.docx_path)
        print("DOCX saved successfully.")
        
        print(f"Saving Markdown to {self.md_path}...")
        with open(self.md_path, "w", encoding="utf-8") as f:
            f.write("\n".join(self.md_lines))
        print("Markdown saved successfully.")
